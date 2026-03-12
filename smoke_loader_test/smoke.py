import tkinter as tk
from tkinter import filedialog, messagebox
import re
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill


# ===== функция чтения файла с автоопределением кодировки =====

def read_log_file(path):
    encodings = ["utf-8", "cp1251", "windows-1251"]

    for enc in encodings:
        try:
            with open(path, encoding=enc) as f:
                return f.readlines()
        except UnicodeDecodeError:
            continue

    raise Exception("Не удалось определить кодировку лог файла")


# ===== GUI =====

root = tk.Tk()
root.withdraw()


# ===== выбор лог файла =====

log_file = filedialog.askopenfilename(
    title="Выберите лог файл loader",
    filetypes=[("Log files", "*.log"), ("All files", "*.*")]
)

if not log_file:
    messagebox.showerror("Ошибка", "Лог файл не выбран")
    exit()


# ===== имя отчета =====

now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
default_name = f"smoke_report_{now}.xlsx"


# ===== выбор места сохранения =====

report_file = filedialog.asksaveasfilename(
    title="Куда сохранить отчет",
    defaultextension=".xlsx",
    initialfile=default_name,
    filetypes=[("Excel files", "*.xlsx")]
)

if not report_file:
    messagebox.showerror("Ошибка", "Не выбрано место сохранения отчета")
    exit()


# ===== читаем лог =====

lines = read_log_file(log_file)


# ===== парсинг =====

tests = []
current_file = None
expected = None

for line in lines:

    m = re.search(r'Processing file (\S+)', line)
    if m:
        current_file = m.group(1)
        expected = None

    m = re.search(r'comment:\s*(.+)', line)
    if m:
        expected = m.group(1).strip()

    m = re.search(r'Finished file (\S+); status:\s*(\w+)', line)
    if m:
        file_name = m.group(1)
        status = m.group(2)

        tests.append((file_name, expected, status))


# ===== анализ =====

rows = []
failed_files = []

pass_count = 0
fail_count = 0

for file_name, expected, status in tests:

    if expected is None:
        continue

    if "OK" in expected and status == "done":
        result = "PASS"
        pass_count += 1

    elif "ERROR" in expected and status == "rejected":
        result = "PASS"
        pass_count += 1

    else:
        result = "FAIL"
        fail_count += 1
        failed_files.append(file_name)

    rows.append([file_name, expected, status, result])


# ===== итог =====

total = pass_count + fail_count
percent = (pass_count / total) * 100 if total else 0


# ===== Excel отчет =====

wb = Workbook()
ws = wb.active
ws.title = "Smoke Report"


# цвета
green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


# заголовок
ws.append(["file", "expected", "actual_status", "result"])


# строки
for r in rows:

    ws.append(r)

    if r[3] == "PASS":
        for cell in ws[ws.max_row]:
            cell.fill = green
    else:
        for cell in ws[ws.max_row]:
            cell.fill = red


# автоподбор ширины колонок
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter

    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[column].width = max_length + 3


# сохраняем
wb.save(report_file)


# ===== вывод результата =====

summary = f"""
SMOKE TEST RESULT

PASS: {pass_count}
FAIL: {fail_count}
TOTAL: {total}

SUCCESS RATE: {percent:.2f} %
"""

if failed_files:

    summary += "\nFAILED FILES:\n"

    for f in failed_files:
        summary += f"{f}\n"

summary += f"\n\nОтчет сохранен:\n{report_file}"

messagebox.showinfo("Smoke тест завершен", summary)


# ===== открыть папку =====

report_folder = os.path.dirname(report_file)
os.startfile(report_folder)